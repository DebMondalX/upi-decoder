"""Human review: spreadsheet out, verified labels back in.

  Train pool (teacher pre-filled, you correct):
    python scripts/04_review.py export --split train_pool
  Test sets (blank, you label from scratch, so teacher bias doesn't leak into the eval):
    python scripts/04_review.py export --split test_in --blank
    python scripts/04_review.py export --split test_unseen --blank
  Bring them back:
    python scripts/04_review.py import --split train_pool
  Synthetic dry run (gold labels instead of a human):
    python scripts/04_review.py from-gold --split train_pool --gold data/synthetic/gold.jsonl

The .xlsx has dropdowns for category / counterparty_type and is sorted so similar
lines sit together: you mostly scan down and fix the odd row. Put 'skip' in the
status column to drop a row (unclear, junk). Flags are comma-separated.
"""

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from upi_decoder.io import read_jsonl, write_jsonl  # noqa: E402
from upi_decoder.schema import CATEGORIES, COUNTERPARTY_TYPES, FLAGS, validate  # noqa: E402

COLS = ["txn_id", "bank", "date", "direction", "amount", "narration",
        "merchant", "counterparty_type", "category", "flags", "status", "notes"]


def export(args):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    rows = read_jsonl(f"data/splits/{args.split}.jsonl")
    teacher = {} if args.blank else {t["txn_id"]: t for t in read_jsonl(args.teacher) if t.get("label")}
    rows.sort(key=lambda r: (r.get("dedup_key", ""), r["date"]))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "review"
    ws.append(COLS)
    missing = 0
    for r in rows:
        lab = teacher.get(r["txn_id"], {}).get("label") or {}
        if not args.blank and not lab:
            missing += 1
        ws.append([r["txn_id"], r["bank"], r["date"], r["direction"], r["amount"], r["narration"],
                   lab.get("merchant", ""), lab.get("counterparty_type", ""), lab.get("category", ""),
                   ", ".join(lab.get("flags", [])), "", ""])

    n = len(rows) + 1
    for col, opts in (("H", COUNTERPARTY_TYPES), ("I", list(CATEGORIES))):
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n}")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
    widths = {"A": 14, "B": 7, "C": 11, "D": 8, "E": 10, "F": 70, "G": 24, "H": 16, "I": 16, "J": 18, "K": 8, "L": 20}
    for c, w in widths.items():
        ws.column_dimensions[c].width = w
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        row[0].alignment = Alignment(wrap_text=False)
    ws.freeze_panes = "G2"

    ref = wb.create_sheet("reference")
    ref.append(["category", "meaning"])
    for k, v in CATEGORIES.items():
        ref.append([k, v])
    ref.append([])
    ref.append(["flags (comma-separated)", ", ".join(FLAGS)])
    ref.append(["counterparty_type", ", ".join(COUNTERPARTY_TYPES)])
    ref.column_dimensions["A"].width = 22
    ref.column_dimensions["B"].width = 70

    out = Path(f"data/review/{args.split}.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out} ({len(rows)} rows{'' if args.blank else f', {missing} without a teacher label'})")


def _import(args):
    import openpyxl

    path = Path(args.file or f"data/review/{args.split}.xlsx")
    ws = openpyxl.load_workbook(path, read_only=True)["review"]
    it = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(it)]
    split_rows = {r["txn_id"]: r for r in read_jsonl(f"data/splits/{args.split}.jsonl")}

    out, errors, skipped, unlabelled = [], [], 0, 0
    for i, vals in enumerate(it, start=2):
        d = dict(zip(header, vals))
        if not d.get("txn_id"):
            continue
        if str(d.get("status") or "").strip().lower() == "skip":
            skipped += 1
            continue
        if not d.get("category"):
            unlabelled += 1
            continue
        lab, errs = validate({
            "merchant": str(d.get("merchant") or "").strip(),
            "counterparty_type": str(d.get("counterparty_type") or "").strip(),
            "category": str(d.get("category") or "").strip(),
            "flags": [f.strip() for f in str(d.get("flags") or "").split(",") if f.strip()],
        })
        if errs:
            errors.append(f"row {i}: {'; '.join(errs)}")
            continue
        base = split_rows.get(str(d["txn_id"]))
        if base is None:
            errors.append(f"row {i}: txn_id not in split")
            continue
        out.append({**base, "label": lab.to_dict()})

    for e in errors[:30]:
        print("  ", e)
    if errors:
        print(f"{len(errors)} rows with errors (not imported). Fix them in the sheet and re-run.")
    dest = Path(f"data/labeled/{args.split}.jsonl")
    write_jsonl(dest, out)
    print(f"imported {len(out)} labels -> {dest}  (skipped {skipped}, unlabelled {unlabelled})")


def from_gold(args):
    gold = {g["txn_id"]: g["label"] for g in read_jsonl(args.gold)}
    rows = read_jsonl(f"data/splits/{args.split}.jsonl")
    out = [{**r, "label": gold[r["txn_id"]]} for r in rows if r["txn_id"] in gold]
    write_jsonl(f"data/labeled/{args.split}.jsonl", out)
    print(f"{args.split}: {len(out)}/{len(rows)} rows labelled from gold")


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("export")
    e.add_argument("--split", required=True)
    e.add_argument("--blank", action="store_true")
    e.add_argument("--teacher", default="data/interim/teacher_labels.jsonl")
    i = sub.add_parser("import")
    i.add_argument("--split", required=True)
    i.add_argument("--file", default=None)
    g = sub.add_parser("from-gold")
    g.add_argument("--split", required=True)
    g.add_argument("--gold", required=True)
    args = ap.parse_args()
    {"export": export, "import": _import, "from-gold": from_gold}[args.cmd](args)


if __name__ == "__main__":
    main()
